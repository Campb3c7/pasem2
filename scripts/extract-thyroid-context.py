"""Preserve thyroid lecture context without committing the oversized PPTX files.

Usage: python scripts/extract-thyroid-context.py <directory containing both decks>
Extracts every slide's editable text, tables, notes, hyperlinks, alt text,
and embedded media. Original media bytes are preserved (no recompression).
"""
from pathlib import Path
from zipfile import ZipFile
import hashlib
import json
import sys
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'source-material/endocrine'
DECKS = [('Thyroid disease - students 2026.pptx', 'thyroid-disease'),
         ('Thyroid Pharmacology - students.pptx', 'thyroid-pharmacology')]

def extract(source, slug):
    deck = Presentation(source)
    media_dir = OUT / 'assets' / slug
    media_dir.mkdir(parents=True, exist_ok=True)
    manifest = []
    with ZipFile(source) as archive:
        for member in archive.namelist():
            if member.startswith('ppt/media/') and not member.endswith('/'):
                blob = archive.read(member)
                target = media_dir / Path(member).name
                target.write_bytes(blob)
                manifest.append({'file': target.name, 'bytes': len(blob),
                                 'sha256': hashlib.sha256(blob).hexdigest()})
    lines = [f'# {deck.slides[0].shapes.title.text if deck.slides[0].shapes.title else slug}',
             '', f'Source: `{source.name}` — {len(deck.slides)} slides.',
             '', 'Complete slide-by-slide reference extraction. Text and table cells are preserved, '
             'including lecture wording. Embedded visuals are linked below and retained at original quality. '
             'Read those images when a slide is primarily a diagram or image-based table. '
             'This is lecture source context, not independently updated clinical guidance.', '']
    counts = {'slides': len(deck.slides), 'text_blocks': 0, 'tables': 0, 'pictures': 0, 'notes': 0}
    def shape_content(shape):
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for child in shape.shapes:
                shape_content(child)
            return
        if shape.has_text_frame:
            paragraphs = [p for p in shape.text_frame.paragraphs if p.text.strip()]
            if paragraphs:
                counts['text_blocks'] += 1
                lines.append('')
                for p in paragraphs:
                    lines.append('  ' * p.level + '- ' + p.text.replace('\n', '\n  '))
                    for run in p.runs:
                        if run.hyperlink.address:
                            lines.append('  Link: ' + run.hyperlink.address)
        if shape.has_table:
            counts['tables'] += 1
            lines.extend(['', 'Table (all rows and cells, including merged-cell origins):'])
            for i, row in enumerate(shape.table.rows, 1):
                lines.append(f'- Row {i}: ' + ' | '.join(cell.text.replace('\n', ' / ') for cell in row.cells))
        for prop in shape._element.xpath('.//p:cNvPr'):
            for name in ('descr', 'title'):
                if prop.get(name):
                    lines.append(f'\nVisual {name}: {prop.get(name)}')
        for blip in shape._element.xpath('.//a:blip'):
            rid = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
            linked = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}link')
            if rid:
                part = shape.part.related_part(rid)
                name = Path(str(part.partname)).name
                counts['pictures'] += 1
                lines.extend(['', f'![Slide visual: {shape.name}](../assets/{slug}/{name})',
                              f'Visual location: x={shape.left}, y={shape.top}, width={shape.width}, height={shape.height} EMU.'])
            elif linked:
                lines.append('\nExternal visual (not embedded): ' + shape.part.rels[linked].target_ref)
        if getattr(shape, 'has_chart', False):
            lines.append('\nChart XML (retains chart labels and cached values):\n```xml\n' +
                         etree.tostring(shape.chart._chartSpace, encoding='unicode') + '\n```')
    for n, slide in enumerate(deck.slides, 1):
        lines.extend(['', f'## Slide {n}', ''])
        for shape in slide.shapes:
            shape_content(shape)
        if slide.has_notes_slide:
            text = slide.notes_slide.notes_text_frame.text
            if text.strip():
                counts['notes'] += 1
                lines.extend(['', '### Speaker notes', '', text])
        # Retain relationship targets for embedded objects, linked diagrams, and media.
        targets = []
        for r in slide.part.rels.values():
            if r.reltype.endswith(('/slideLayout', '/notesSlide', '/image')):
                continue
            if r.is_external:
                targets.append(r.target_ref)
                continue
            part = r.target_part
            name = Path(str(part.partname)).name
            if '/media/' in str(part.partname):
                targets.append(f'[Embedded media](../assets/{slug}/{name})')
            else:
                resource_dir = OUT / 'assets' / slug / 'resources'
                resource_dir.mkdir(exist_ok=True)
                target = resource_dir / name
                target.write_bytes(part.blob)
                targets.append(f'[Slide resource](../assets/{slug}/resources/{name})')
                if name.endswith('.xml'):
                    xml = etree.fromstring(part.blob)
                    texts = xml.xpath('//*[local-name()="t"]/text()')
                    if texts:
                        lines.extend(['', 'Diagram/resource text:'] + ['- ' + text for text in texts])
        if targets:
            lines.extend(['', 'Additional slide resources:'] + ['- ' + t for t in targets])
    text_dir = OUT / 'text'
    text_dir.mkdir(exist_ok=True)
    (text_dir / (slug + '.md')).write_text('\n'.join(lines) + '\n', encoding='utf-8')
    record = {'source': source.name, 'source_sha256': hashlib.sha256(source.read_bytes()).hexdigest(),
              'counts': counts, 'media': manifest}
    (media_dir / 'manifest.json').write_text(json.dumps(record, indent=2) + '\n', encoding='utf-8')
    print(slug, json.dumps(counts), 'media files:', len(manifest))

if __name__ == '__main__':
    for filename, slug in DECKS:
        extract(Path(sys.argv[1]) / filename, slug)
