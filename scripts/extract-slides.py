#!/usr/bin/env python3
"""Extract slide text + speaker notes from lecture decks into searchable Markdown.

Usage: python3 scripts/extract-slides.py source-material/<course>
Reads <course>/slides/*.pptx and writes <course>/text/<deck>.md.

Text, tables, and notes are captured. Pictures are not; slides that carry
content in pictures are marked "[picture]". To see those, render the deck:
  soffice --headless --convert-to pdf --outdir <scratch dir> <deck>.pptx
and read the needed PDF pages.
"""
import glob, os, sys
import pptx
from pptx.enum.shapes import MSO_SHAPE_TYPE

def shape_lines(shape, deck):
    lines = []
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for child in shape.shapes: lines += shape_lines(child, deck)
        return lines
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        try: shape.image; embedded = True
        except ValueError: embedded = False
        big = shape.width * shape.height >= 0.08 * deck.slide_width * deck.slide_height
        if big: lines.append("- [picture]" if embedded else "- [picture: linked, not embedded in deck]")
    if shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            text = "".join(run.text for run in para.runs).strip()
            if text: lines.append("  " * para.level + "- " + text)
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            lines.append("| " + " | ".join(cell.text.replace("\n", " ").strip() for cell in row.cells) + " |")
    return lines

def extract(src, dst):
    deck = pptx.Presentation(src)
    out = []
    for number, slide in enumerate(deck.slides, 1):
        out.append("\n## Slide %d\n" % number)
        for shape in slide.shapes: out += shape_lines(shape, deck)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes: out.append("\n> Notes: " + notes.replace("\n", "\n> "))
    with open(dst, "w") as handle: handle.write("\n".join(out) + "\n")

course = sys.argv[1]
os.makedirs(os.path.join(course, "text"), exist_ok=True)
for src in sorted(glob.glob(os.path.join(course, "slides", "*.pptx"))):
    name = os.path.splitext(os.path.basename(src))[0]
    extract(src, os.path.join(course, "text", name + ".md"))
    print("extracted", name)
